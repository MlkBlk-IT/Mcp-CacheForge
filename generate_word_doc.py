from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

doc = Document()

# Titre
title = doc.add_heading('CacheForge : Proxy LLM avec Cache Sémantique Local', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Métadonnées
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Statut : ').bold = True
p.add_run('En développement (Résolution bug métriques)\n')
p.add_run('Dernière mise à jour : ').bold = True
p.add_run('2026-10-22 03:00 EST\n')
p.add_run('Auteur : ').bold = True
p.add_run('Malik (DEP Soutien Informatique, EMICA)')

# 1. Problématique
doc.add_heading('1. Problématique Résolue', level=1)
doc.add_paragraph(
    "L'utilisation de modèles de langage locaux (via LM Studio) pour des tâches IT récurrentes "
    "présente trois goulots d'étranglement :\n"
    "1. Latence élevée : La génération de réponses identiques consomme inutilement du temps CPU/RAM.\n"
    "2. Absence de cache sémantique : Les systèmes standards hachent le prompt brut. 'Explique Hyper-V' "
    "et 'S'il te plaît, explique Hyper-V stp' génèrent deux requêtes distinctes.\n"
    "3. Manque d'observabilité : Impossible de mesurer les tokens ou le temps de calcul économisés."
)

# 2. Solution
doc.add_heading('2. Solution : Architecture CacheForge', level=1)
doc.add_paragraph(
    "CacheForge est un proxy HTTP asynchrone (FastAPI) qui s'intercale entre le client et le backend LLM (LM Studio :1234)."
)

p_workflow = doc.add_paragraph()
p_workflow.add_run('Pipeline de traitement :').bold = True
doc.add_paragraph('1. Réception : POST /v1/chat/completions sur le port 8080.', style='List Bullet')
doc.add_paragraph('2. Normalisation : NFKD + Lowercase + Aho-Corasick (suppression de 1114 motifs de bruit comme "svp", "stp").', style='List Bullet')
doc.add_paragraph('3. Hachage : SHA256 du payload normalisé (model, messages, temperature, tools).', style='List Bullet')
doc.add_paragraph('4. Décision : HIT (retour instantané + métriques) ou MISS (forward vers LLM + sauvegarde SQLite avec TTL).', style='List Bullet')

doc.add_heading('Stack Technique', level=2)
doc.add_paragraph('Backend : Python 3.11, FastAPI, uvicorn, httpx', style='List Bullet')
doc.add_paragraph('Cache : SQLite3 (rotation TTL 24h / max 10 000 entrées)', style='List Bullet')
doc.add_paragraph('Normalisation : pyahocorasick, nltk (Snowball Stemmer)', style='List Bullet')
doc.add_paragraph('Infrastructure : pythonw.exe (headless), pystray (System Tray), Task Scheduler Windows', style='List Bullet')
doc.add_paragraph('Monitoring : Dashboard web local (http://localhost:8080/dashboard)', style='List Bullet')

# 3. Journal de Développement
doc.add_heading('3. Journal de Développement (Timestamped)', level=1)

table = doc.add_table(rows=1, cols=3)
table.style = 'Light Shading Accent 1'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Date / Heure (EST)'
hdr_cells[1].text = 'Phase / Action'
hdr_cells[2].text = 'Résultat'

data = [
    ('2026-10-21 19:00', 'Initialisation du proxy basique + Forward vers LM Studio.', 'Fonctionnel (Miss uniquement)'),
    ('2026-10-21 20:30', 'Implémentation SQLite + Logique HIT/MISS basique.', 'Fonctionnel, mais hachage sur texte brut'),
    ('2026-10-22 00:15', 'Correction régression SQLite + Suppression middleware Auth (Keyless).', 'Stable'),
    ('2026-10-22 00:45', 'Déploiement Infrastructure Headless : tray.py, Task Scheduler, Dashboard.', 'Fonctionnel'),
    ('2026-10-22 01:15', 'Réécriture NormalizationPipeline (Aho-Corasick + NFKD).', 'Validé : "test" == "SVP test stp"'),
    ('2026-10-22 02:15', 'Tentative d\'ajout des métriques avancées (saved_tokens, gen_time_ms).', 'ÉCHEC : Erreur 500 Internal Server Error')
]

for date, action, result in data:
    row_cells = table.add_row().cells
    row_cells[0].text = date
    row_cells[1].text = action
    row_cells[2].text = result

# 4. État Actuel
doc.add_heading('4. État Actuel & Prochaines Étapes', level=1)

p_bug = doc.add_paragraph()
p_bug.add_run('Bug Bloquant Actuel : ').bold = True
p_bug.add_run(
    "Depuis 02:15 EST, toutes les requêtes retournent une Erreur 500. "
    "Le script apply_final_patch.py a probablement corrompu l'injection des variables (forward_start, saved_usage) "
    "ou la structure du cache_payload dans server.py."
)

p_next = doc.add_paragraph()
p_next.add_run('Plan d\'action :').bold = True
doc.add_paragraph('1. Isoler le crash en lançant tray.py en mode console pour obtenir le traceback Python exact.', style='List Number')
doc.add_paragraph('2. Rollback : Restaurer server.py.bak14 (dernière version stable).', style='List Number')
doc.add_paragraph('3. Patch chirurgical : Réécrire l\'injection des métriques ligne par ligne.', style='List Number')
doc.add_paragraph('4. Validation : Exécuter la suite de 5 tests (scénario Hyper-V DHCP) pour valider les métriques.', style='List Number')

# Sauvegarde
output_path = Path(r"D:\CacheForge\CacheForge_Project_Documentation.docx")
doc.save(str(output_path))
print(f"✅ Document Word généré avec succès : {output_path}")
